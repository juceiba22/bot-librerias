"""Carga en la base los listados de la distribuidora de SBS (Excel y PDF) como canales aparte.

  novedades   -> canal 'distribuidora' (novedades mensuales, ISBN / titulo / autor / PVP / sello)
  textos      -> canal 'textos'        (novedades de textos escolares)
  elt         -> canal 'elt'           (ingles para institutos, PDF)

Los archivos son listas internas del cliente: NO se copian al repositorio, se pasan por ruta.

Uso:
  python cargar_sbs_distribuidora.py --novedades SBS_NOVEDADES.xlsx --textos SBS_NOVEDADES_TEXTO.xlsx --elt SBS_NOVEDADES_ELT.pdf [--negocio sbs]
"""
import argparse
import json
import os
import re
import urllib.request

import openpyxl
from pypdf import PdfReader

RUNNER = os.environ.get("SQL_RUNNER_URL", "")  # URL del webhook utilitario de SQL (no se guarda en el repo)
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


def limpio(v):
    if v is None:
        return None
    t = re.sub(r"\s+", " ", str(v)).strip()
    return t if t and t.upper() not in ("NONE", "NAN") else None


def isbn_de(v):
    d = re.sub(r"\D", "", str(v).split(".")[0]) if v is not None else ""
    return d if len(d) in (10, 13) else None


def precio_de(v):
    try:
        p = float(v)
    except (TypeError, ValueError):
        return None
    return p if p > 0 else None


def novedades(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    por_isbn = {}
    for ws in wb.worksheets:
        mes = ws.title.strip().upper()
        orden = MESES.index(mes) if mes in MESES else -1
        for row in ws.iter_rows(min_row=5, values_only=True):
            if not row or len(row) < 6:
                continue
            isbn, titulo = isbn_de(row[1]), limpio(row[2])
            if not isbn or not titulo:
                continue
            fila = {
                "canal": "distribuidora", "grupo": "Novedades de la distribuidora",
                "categoria": f"Novedades {mes.title()}", "sku": isbn, "isbn": isbn,
                "nombre": titulo, "marca": limpio(row[3]), "editorial": limpio(row[5]),
                "precio": precio_de(row[4]),
                "descripcion": f"Novedad de {mes.title()}. Precio de lista (PVP), no es precio mayorista." + (f" Sello: {limpio(row[5])}." if limpio(row[5]) else ""),
                "_orden": orden,
            }
            if isbn not in por_isbn or orden >= por_isbn[isbn]["_orden"]:
                por_isbn[isbn] = fila
    return list(por_isbn.values())


def textos(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out, vistos = [], set()
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if not row or len(row) < 4:
            continue
        isbn, titulo = isbn_de(row[0]), limpio(row[1])
        if not isbn or not titulo or isbn in vistos:
            continue
        vistos.add(isbn)
        ed = limpio(row[3])
        out.append({
            "canal": "textos", "grupo": "Textos escolares", "categoria": (ed or "Textos escolares").title(),
            "sku": isbn, "isbn": isbn, "nombre": titulo, "marca": None, "editorial": ed,
            "precio": precio_de(row[2]),
            "descripcion": "Texto escolar, novedad 2026." + (f" Editorial: {ed.title()}." if ed else ""),
        })
    return out


PREFIJOS_EDITORIAL = [("9781009", "Cambridge University Press"), ("9781292", "Pearson"), ("9788543", "Pearson")]


def elt(path):
    out, vistos, serie = [], set(), None
    for pag in PdfReader(path).pages:
        for linea in (pag.extract_text() or "").split("\n"):
            l = linea.strip()
            m = re.match(r"^(\d{13})\s+(.*)$", l)
            if m:
                isbn, resto = m.group(1), m.group(2)
                pr = re.search(r"\$\s*([\d.]+)\s*$", resto)
                precio = float(pr.group(1).replace(".", "")) if pr else None
                nombre = limpio(resto[:pr.start()] if pr else resto)
                if not nombre or isbn in vistos:
                    continue
                vistos.add(isbn)
                ed = next((e for p, e in PREFIJOS_EDITORIAL if isbn.startswith(p)), None)
                out.append({
                    "canal": "elt", "grupo": "ELT - ingles para institutos", "categoria": serie or "ELT",
                    "sku": isbn, "isbn": isbn, "nombre": nombre, "marca": None, "editorial": ed,
                    "precio": precio,
                    "descripcion": "Material de enseñanza de inglés (ELT), novedad 2027." + (f" Serie: {serie}." if serie else "")
                                   + (f" Editorial: {ed}." if ed else "") + ("" if precio else " Precio a consultar."),
                })
            elif l and not l.upper().startswith(("NOVEDADES", "PEARSON", "CAMBRIDGE")):
                s = re.sub(r"\s*(New\s*[Ee]dition|New)\s*!?\s*$", "", l, flags=re.I).strip()
                if s:
                    serie = re.sub(r"(\d)(St|Nd|Rd|Th)(?![a-z])", lambda m: m.group(1) + m.group(2).lower(), s.title())
    return out


def q(v):
    return "NULL" if v is None else "'" + str(v).replace("\x00", "").replace("'", "''") + "'"


def cargar(filas, slug):
    sub = f"(SELECT id FROM negocios WHERE slug = {q(slug)})"
    for i in range(0, len(filas), 120):
        lote = filas[i:i + 120]
        valores = ",\n".join(
            f"({sub}, {q(f['canal'])}, {q(f['grupo'])}, {q(f['categoria'])}, {q(f['sku'])}, {q(f['isbn'])}, {q(f['nombre'])}, "
            f"{q(f['marca'])}, {q(f['editorial'])}, {f['precio'] if f['precio'] is not None else 'NULL'}, 'ARS', 'Novedad', {q(f['descripcion'])})"
            for f in lote
        )
        sql = ("INSERT INTO productos (negocio_id, canal, grupo, segmento_demo, sku, isbn, nombre, marca, editorial, precio, moneda, disponibilidad, descripcion) VALUES\n"
               + valores + "\nON CONFLICT (negocio_id, canal, sku) DO UPDATE SET nombre = EXCLUDED.nombre, marca = EXCLUDED.marca, "
               "editorial = EXCLUDED.editorial, precio = EXCLUDED.precio, grupo = EXCLUDED.grupo, segmento_demo = EXCLUDED.segmento_demo, "
               "descripcion = EXCLUDED.descripcion, actualizado_en = now();")
        req = urllib.request.Request(RUNNER, data=json.dumps({"sql": sql}).encode("utf-8"), headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=90) as r:
            resp = r.read().decode("utf-8")[:60]
        print(f"    lote {i // 120 + 1}: {len(lote)} filas -> {resp}")


def main():
    if not RUNNER:
        raise SystemExit('Falta la variable de entorno SQL_RUNNER_URL')
    ap = argparse.ArgumentParser()
    ap.add_argument("--novedades"); ap.add_argument("--textos"); ap.add_argument("--elt")
    ap.add_argument("--negocio", default="sbs")
    a = ap.parse_args()
    bloques = []
    if a.novedades: bloques.append(("distribuidora", novedades(a.novedades)))
    if a.textos: bloques.append(("textos", textos(a.textos)))
    if a.elt: bloques.append(("elt", elt(a.elt)))
    for canal, filas in bloques:
        con_precio = sum(1 for f in filas if f["precio"] is not None)
        print(f"{canal}: {len(filas)} productos ({con_precio} con precio)")
        borrar = f"DELETE FROM productos WHERE canal = {q(canal)} AND negocio_id = (SELECT id FROM negocios WHERE slug = {q(a.negocio)});"
        req = urllib.request.Request(RUNNER, data=json.dumps({"sql": borrar}).encode("utf-8"), headers={"Content-Type": "application/json"})
        urllib.request.urlopen(req, timeout=60).read()
        cargar(filas, a.negocio)


if __name__ == "__main__":
    main()
